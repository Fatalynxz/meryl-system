import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment
import os

excel_file = "MERYL_SHOES_GANTT_CHART_AND_PERT_CPM.xlsx"
pert_img_path = os.path.join("diagrams_output", "MERYL_SHOES_PERT_CHART.png")
cpm_img_path = os.path.join("diagrams_output", "MERYL_SHOES_CPM_NETWORK_DIAGRAM.png")

wb = openpyxl.load_workbook(excel_file)

# 1. PERT Sheet
pert_sheet_name = "PERT Network Diagram"
if pert_sheet_name in wb.sheetnames:
    del wb[pert_sheet_name]
ws_pert = wb.create_sheet(title=pert_sheet_name)
ws_pert.views.sheetView[0].showGridLines = True

ws_pert["A1"] = "MERYL SHOES SYSTEM — PERT NETWORK DIAGRAM (APPENDIX I)"
ws_pert["A1"].font = Font(name="Calibri", size=14, bold=True, color="1B5E20")

ws_pert["A2"] = "Critical Path Duration: 159 Calendar Days (May 5 – Oct 10, 2026) | Chain: A → B → C → D → F → G → H → I → K → L → M/N → P → Q → S → T/U → V → W → X"
ws_pert["A2"].font = Font(name="Calibri", size=10, italic=True, color="4B5563")

img_pert = Image(pert_img_path)
# Scale image to comfortable spreadsheet viewing dimensions (1400 x 875 px)
img_pert.width = 1400
img_pert.height = 875
ws_pert.add_image(img_pert, "B4")


# 2. CPM Sheet
cpm_sheet_name = "CPM Network Diagram"
if cpm_sheet_name in wb.sheetnames:
    del wb[cpm_sheet_name]
ws_cpm = wb.create_sheet(title=cpm_sheet_name)
ws_cpm.views.sheetView[0].showGridLines = True

ws_cpm["A1"] = "MERYL SHOES SYSTEM — CRITICAL PATH METHOD (CPM) NETWORK DIAGRAM"
ws_cpm["A1"].font = Font(name="Calibri", size=14, bold=True, color="1B5E20")

ws_cpm["A2"] = "Detailed ES, EF, LS, LF, Duration (D), and Float/Slack [S] per Activity Node"
ws_cpm["A2"].font = Font(name="Calibri", size=10, italic=True, color="4B5563")

img_cpm = Image(cpm_img_path)
img_cpm.width = 1400
img_cpm.height = 875
ws_cpm.add_image(img_cpm, "B4")

wb.save(excel_file)
print(f"Successfully embedded PERT and CPM network diagrams into '{excel_file}'.")
print(f"Updated Sheets: {wb.sheetnames}")

